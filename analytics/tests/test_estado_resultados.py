"""Tests for reading the accountant's P&L statement and filling the sheet (US43)."""
import datetime
from decimal import Decimal
import pytest
from openpyxl import Workbook, load_workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from analytics.estado_resultados import (
    CREDIT_ROWS,
    FORMULA_ROWS,
    ROW_ACCOUNTS,
    StatementError,
    double_mapped_accounts,
    merge,
    read_statement,
    reconcile,
)
from analytics.unified_excel import UnifiedUpdateError, update_estados_resultados

# Right edge of each monthly column, matching the real report's layout.
_EDGES = (297, 393, 489, 585)


def _statement_pdf(path, months=("2026/01", "2026/02"), lines=None):
    """A minimal 'ESTADO DE RESULTADOS' with right-aligned monthly columns."""
    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.setFont("Helvetica", 7)
    y = 720
    header = "CUENTA D E S C R I P C I O N " + " ".join(
        f"NETOS DESDE {m}" for m in months
    )
    pdf.drawString(21, y, header)
    y -= 14
    for label, amounts in lines or []:
        pdf.drawString(21, y, label)
        for index, amount in enumerate(amounts):
            if amount is None:
                continue
            pdf.drawRightString(_EDGES[index], y, amount)
        y -= 14
    pdf.save()
    return path


def _bold(text):
    """How the report renders a bold subtotal: every glyph printed twice."""
    return "".join(c * 2 for c in text)


def _simple(path):
    """One sales account, one expense account and the totals that bound them."""
    return _statement_pdf(
        path,
        lines=[
            ("41401501 COMESTIBLES", ["1,000.00CR", "2,000.00CR"]),
            (_bold("**** OPERACIONALES"), [_bold("1,000.00CR"), _bold("2,000.00CR")]),
            ("520506 SUELDOS", ["300.00", "400.00"]),
            (_bold("** GASTOS"), [_bold("300.00"), _bold("400.00")]),
            ("61401501 COMESTIBLES", ["100.00", "150.00"]),
            (_bold("** COSTOS"), [_bold("100.00"), _bold("150.00")]),
        ],
    )


@pytest.fixture
def statement(tmp_path):
    return read_statement(_simple(tmp_path / "ef.pdf"))


def test_reads_the_months_from_the_header(statement):
    assert statement.months == ["2026-01", "2026-02"]


def test_reads_account_amounts_per_month(statement):
    assert statement.amounts[("520506", "2026-01")] == Decimal("300")
    assert statement.amounts[("520506", "2026-02")] == Decimal("400")


def test_income_is_credit_and_the_sheet_keeps_it_positive(statement):
    # The statement prints sales as "CR"; the sheet's VENTAS row is positive.
    assert statement.amounts[("41401501", "2026-01")] == Decimal("-1000")
    assert statement.row_value(3, "2026-01") == Decimal("1000")
    assert 3 in CREDIT_ROWS


def test_amounts_land_on_the_month_their_column_says(tmp_path):
    # A row that only moved in the second month prints one amount, on the right.
    path = _statement_pdf(
        tmp_path / "gap.pdf",
        lines=[("520506 SUELDOS", [None, "999.00"])],
    )
    statement = read_statement(path)
    assert ("520506", "2026-01") not in statement.amounts
    assert statement.amounts[("520506", "2026-02")] == Decimal("999")


def test_bold_subtotals_are_undoubled(statement):
    assert statement.subtotals[("** GASTOS", "2026-01")] == Decimal("300")


def test_group_lines_are_not_treated_as_accounts(tmp_path):
    path = _statement_pdf(
        tmp_path / "groups.pdf",
        lines=[
            ("5205 GASTOS DE PERSONAL", []),
            ("520506 SUELDOS", ["300.00", "400.00"]),
        ],
    )
    statement = read_statement(path)
    assert "5205" not in statement.leaf_accounts
    assert "520506" in statement.leaf_accounts


def test_a_file_without_the_period_header_is_rejected(tmp_path):
    pdf = canvas.Canvas(str(tmp_path / "other.pdf"), pagesize=letter)
    pdf.drawString(21, 700, "BALANCE GENERAL")
    pdf.save()
    with pytest.raises(StatementError):
        read_statement(tmp_path / "other.pdf")


def test_merge_keeps_every_month_in_order(tmp_path):
    first = read_statement(_simple(tmp_path / "a.pdf"))
    second = read_statement(
        _statement_pdf(
            tmp_path / "b.pdf",
            months=("2026/03", "2026/04"),
            lines=[("520506 SUELDOS", ["500.00", "600.00"])],
        )
    )
    merged = merge([second, first])
    assert merged.months == ["2026-01", "2026-02", "2026-03", "2026-04"]
    assert merged.amounts[("520506", "2026-03")] == Decimal("500")


def test_no_account_is_claimed_by_two_rows():
    assert double_mapped_accounts() == []


def test_reconcile_accepts_a_statement_that_adds_up(statement):
    assert reconcile(statement) == []


def test_reconcile_reports_an_account_no_row_claims(tmp_path):
    path = _statement_pdf(
        tmp_path / "orphan.pdf",
        lines=[("529999 CUENTA NUEVA", ["50.00", "60.00"])],
    )
    problems = reconcile(read_statement(path))
    assert any("529999" in p for p in problems)


def test_reconcile_reports_a_total_that_does_not_match(tmp_path):
    path = _statement_pdf(
        tmp_path / "bad.pdf",
        lines=[
            ("520506 SUELDOS", ["300.00", "400.00"]),
            (_bold("** GASTOS"), [_bold("999.00"), _bold("400.00")]),
        ],
    )
    problems = reconcile(read_statement(path))
    assert any("TOTAL GASTOS" in p for p in problems)


def _sheet(tmp_path):
    """A workbook shaped like the owner's: months across, rubros down."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Estados de Resultados"
    ws["A1"] = "RUBLO"
    ws.cell(row=1, column=2, value=datetime.datetime(2025, 12, 1))
    labels = {3: "COMIDA", 29: "SUELDOS", 7: "TOTAL VENTAS", 43: "SALARIOS"}
    for row, label in labels.items():
        ws.cell(row=row, column=1, value=label)
    ws.cell(row=3, column=2, value=111)
    ws.cell(row=29, column=2, value=222)
    ws.cell(row=7, column=2, value="=+B3+B4-B6-B5")
    ws.cell(row=43, column=2, value="=SUM(B29:B42)")
    path = tmp_path / "master.xlsx"
    wb.save(path)
    return path


def test_update_appends_a_column_per_month(tmp_path, statement):
    summary = update_estados_resultados(_sheet(tmp_path), statement)

    ws = load_workbook(summary["copy"])["Estados de Resultados"]
    assert summary["written"] == ["2026-01", "2026-02"]
    assert ws.cell(row=1, column=3).value == datetime.datetime(2026, 1, 1)
    assert ws.cell(row=1, column=4).value == datetime.datetime(2026, 2, 1)
    assert ws.cell(row=3, column=3).value == 1000  # VENTAS / COMIDA, sign flipped
    assert ws.cell(row=29, column=4).value == 400  # SUELDOS
    # December is untouched.
    assert ws.cell(row=3, column=2).value == 111


def test_update_replicates_the_sheets_own_formulas(tmp_path, statement):
    summary = update_estados_resultados(_sheet(tmp_path), statement)

    ws = load_workbook(summary["copy"])["Estados de Resultados"]
    assert ws.cell(row=7, column=3).value == "=+C3+C4-C6-C5"
    assert ws.cell(row=43, column=4).value == "=SUM(D29:D42)"
    assert 7 in FORMULA_ROWS and 7 not in ROW_ACCOUNTS


def test_update_leaves_untracked_rubros_blank(tmp_path, statement):
    # The accountant books nothing on "PAN CORTESIA" (row 79): blank, not zero,
    # so the owner can tell "no account" from "no movement".
    summary = update_estados_resultados(_sheet(tmp_path), statement)

    ws = load_workbook(summary["copy"])["Estados de Resultados"]
    assert ws.cell(row=79, column=3).value is None
    assert ws.cell(row=29, column=3).value == 300


def test_update_overwrites_a_month_already_there(tmp_path, statement):
    first = update_estados_resultados(_sheet(tmp_path), statement)
    second = update_estados_resultados(first["copy"], statement)

    assert second["overwritten"] == ["2026-01", "2026-02"]
    assert second["written"] == []
    ws = load_workbook(second["copy"])["Estados de Resultados"]
    assert ws.cell(row=1, column=5).value is None  # no duplicate columns


def test_update_never_touches_the_original(tmp_path, statement):
    path = _sheet(tmp_path)
    before = path.read_bytes()
    summary = update_estados_resultados(path, statement)

    assert summary["copy"] != path
    assert path.read_bytes() == before


def test_update_rejects_a_sheet_without_month_headers(tmp_path, statement):
    wb = Workbook()
    wb.active.title = "Estados de Resultados"
    path = tmp_path / "empty.xlsx"
    wb.save(path)
    with pytest.raises(UnifiedUpdateError):
        update_estados_resultados(path, statement)


def test_update_rejects_a_missing_file(tmp_path, statement):
    with pytest.raises(UnifiedUpdateError):
        update_estados_resultados(tmp_path / "nope.xlsx", statement)


def test_four_columns_map_to_four_months(tmp_path):
    # Guards the column-edge constants: a value at the far right must be the
    # last month, not the first.
    path = _statement_pdf(
        tmp_path / "edges.pdf",
        months=("2026/01", "2026/02", "2026/03", "2026/04"),
        lines=[("520506 SUELDOS", ["1.00", "2.00", "3.00", "4.00"])],
    )
    statement = read_statement(path)
    assert [statement.amounts[("520506", m)] for m in statement.months] == [
        Decimal("1"),
        Decimal("2"),
        Decimal("3"),
        Decimal("4"),
    ]


def test_every_mapped_row_is_outside_the_formula_rows():
    assert not (set(ROW_ACCOUNTS) & FORMULA_ROWS)
