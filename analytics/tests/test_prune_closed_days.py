"""Tests for dropping the closed days from 'Datos totales' (US44)."""
import datetime

import pytest
from openpyxl import Workbook, load_workbook

from analytics.unified_excel import UnifiedUpdateError, prune_closed_days

# 2026-01-05, -12, -19 and -26 are Mondays; -06 is a Tuesday.
_MONDAY = datetime.datetime(2026, 1, 5)


def _sheet(tmp_path, days):
    """A 'Datos totales' sheet: ``days`` is a list of (date, people, bar)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos totales "
    for col, title in enumerate(
        ["Año", "Fecha", "Dia", "Mes", "Día", "Nro Personas", "Nro Cuentas",
         "Tkt Prom/P", "Tkt Prom/C", "Venta Total"], start=1
    ):
        ws.cell(row=1, column=col, value=title)
    for index, (day, people, bar) in enumerate(days):
        row = index + 2
        ws.cell(row=row, column=1, value=day.year)
        ws.cell(row=row, column=2, value=day)
        ws.cell(row=row, column=6, value=people)
        ws.cell(row=row, column=8, value=f'=IFERROR(J{row}/F{row},"")')
        ws.cell(row=row, column=10, value=f"=N{row}+S{row}")
        ws.cell(row=row, column=14, value=bar)
    path = tmp_path / "master.xlsx"
    wb.save(path)
    return path


def _dates(path):
    ws = load_workbook(path)["Datos totales "]
    return [
        ws.cell(row=r, column=2).value.date()
        for r in range(2, ws.max_row + 1)
        if isinstance(ws.cell(row=r, column=2).value, datetime.datetime)
    ]


def test_deletes_the_empty_mondays(tmp_path):
    path = _sheet(tmp_path, [
        (_MONDAY, None, None),                                  # closed Monday
        (datetime.datetime(2026, 1, 6), 30, 500_000),           # Tuesday
        (datetime.datetime(2026, 1, 12), None, None),           # closed Monday
    ])
    summary = prune_closed_days(path)

    assert summary["deleted"] == ["2026-01-05", "2026-01-12"]
    assert _dates(summary["copy"]) == [datetime.date(2026, 1, 6)]


def test_keeps_a_monday_that_had_movement(tmp_path):
    # A holiday Monday: the restaurant did open, so the row stays.
    path = _sheet(tmp_path, [
        (_MONDAY, 25, 900_000),
        (datetime.datetime(2026, 1, 12), None, None),
    ])
    summary = prune_closed_days(path)

    assert summary["deleted"] == ["2026-01-12"]
    assert summary["kept"] == 1
    assert _dates(summary["copy"]) == [datetime.date(2026, 1, 5)]


def test_a_monday_with_sales_but_no_headcount_is_movement(tmp_path):
    path = _sheet(tmp_path, [(_MONDAY, None, 142_591)])
    summary = prune_closed_days(path)

    assert summary["deleted"] == []
    assert _dates(summary["copy"]) == [datetime.date(2026, 1, 5)]


def test_other_weekdays_are_left_alone(tmp_path):
    # An empty Tuesday is a closure we were not asked about, not a Monday.
    path = _sheet(tmp_path, [(datetime.datetime(2026, 1, 6), None, None)])
    summary = prune_closed_days(path)

    assert summary["deleted"] == []
    assert _dates(summary["copy"]) == [datetime.date(2026, 1, 6)]


def test_weekday_none_prunes_every_empty_day(tmp_path):
    path = _sheet(tmp_path, [
        (_MONDAY, None, None),
        (datetime.datetime(2026, 1, 6), None, None),
        (datetime.datetime(2026, 1, 7), 12, 300_000),
    ])
    summary = prune_closed_days(path, weekday=None)

    assert summary["deleted"] == ["2026-01-05", "2026-01-06"]
    assert _dates(summary["copy"]) == [datetime.date(2026, 1, 7)]


def test_surviving_rows_point_their_formulas_at_their_own_row(tmp_path):
    # The kept Tuesday moves from row 3 to row 2; its formulas must follow.
    path = _sheet(tmp_path, [
        (_MONDAY, None, None),
        (datetime.datetime(2026, 1, 6), 30, 500_000),
    ])
    summary = prune_closed_days(path)

    ws = load_workbook(summary["copy"])["Datos totales "]
    assert ws.cell(row=2, column=10).value == "=N2+S2"
    assert ws.cell(row=2, column=8).value == '=IFERROR(J2/F2,"")'


def test_repairs_a_formula_that_already_pointed_elsewhere(tmp_path):
    # The real sheet had rows whose formulas pointed at a neighbour; re-pointing
    # every surviving row fixes those too.
    path = _sheet(tmp_path, [(datetime.datetime(2026, 1, 6), 30, 500_000)])
    wb = load_workbook(path)
    wb["Datos totales "].cell(row=2, column=10, value="=N99+S99")
    wb.save(path)

    summary = prune_closed_days(path)

    ws = load_workbook(summary["copy"])["Datos totales "]
    assert ws.cell(row=2, column=10).value == "=N2+S2"
    assert summary["formulas_repaired"] >= 1


def test_each_rows_own_formula_columns_are_preserved(tmp_path):
    # Rows differ in which columns carry formulas; pruning must not add any.
    path = _sheet(tmp_path, [
        (_MONDAY, None, None),
        (datetime.datetime(2026, 1, 6), 30, 500_000),
    ])
    wb = load_workbook(path)
    wb["Datos totales "].cell(row=3, column=8).value = None  # no Tkt Prom here
    wb.save(path)

    summary = prune_closed_days(path)

    ws = load_workbook(summary["copy"])["Datos totales "]
    assert ws.cell(row=2, column=8).value is None
    assert ws.cell(row=2, column=10).value == "=N2+S2"


def test_never_touches_the_original(tmp_path):
    path = _sheet(tmp_path, [(_MONDAY, None, None)])
    before = path.read_bytes()

    summary = prune_closed_days(path)

    assert summary["copy"] != path
    assert path.read_bytes() == before


def test_rejects_a_missing_file(tmp_path):
    with pytest.raises(UnifiedUpdateError):
        prune_closed_days(tmp_path / "nope.xlsx")


def test_rejects_a_workbook_without_the_sheet(tmp_path):
    wb = Workbook()
    wb.active.title = "Otra cosa"
    path = tmp_path / "wrong.xlsx"
    wb.save(path)
    with pytest.raises(UnifiedUpdateError):
        prune_closed_days(path)


def test_a_closure_typed_as_zero_counts_as_closed(tmp_path):
    # Some closed days were typed as 0 instead of left blank.
    path = _sheet(tmp_path, [(_MONDAY, 0, 0)])
    summary = prune_closed_days(path)

    assert summary["deleted"] == ["2026-01-05"]


def test_zero_on_one_side_only_is_still_movement(tmp_path):
    # No bar sales but people through the door: the place was open.
    path = _sheet(tmp_path, [(_MONDAY, 18, 0)])
    summary = prune_closed_days(path)

    assert summary["deleted"] == []


def test_the_calendar_skeleton_ahead_of_today_is_left_alone(tmp_path):
    # The sheet pre-fills the rest of the year with empty rows; those are not
    # closures, and deleting them would push a future holiday Monday to the
    # bottom of the sheet when its report arrives.
    path = _sheet(tmp_path, [
        (_MONDAY, None, None),
        (datetime.datetime(2026, 12, 28), None, None),
    ])
    summary = prune_closed_days(path, through=datetime.date(2026, 6, 30))

    assert summary["deleted"] == ["2026-01-05"]
    assert _dates(summary["copy"]) == [datetime.date(2026, 12, 28)]
