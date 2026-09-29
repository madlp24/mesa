"""Turnout per operating day, read from the bill-level export (US41)."""
import datetime
from io import StringIO

import pytest
from django.core.management import call_command
from openpyxl import Workbook

from sales.importers.tickets import (
    TicketExportError,
    operating_day,
    parse_closed_at,
    read_daily_covers,
)
from sales.models import DailyCovers


def _export(path, rows):
    """The POS export: a few title lines, then the header, then one row per bill."""
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="REPORTE DE VENTAS")
    for col, label in enumerate(
        ["FOLIO", "CIERRE", "NÚMERODEPERSONAS", "TOTALDEARTÍCULOS"], start=1
    ):
        ws.cell(row=5, column=col, value=label)
    for i, (closed, people) in enumerate(rows, start=6):
        ws.cell(row=i, column=1, value=f"FV{i}")
        ws.cell(row=i, column=2, value=closed)
        ws.cell(row=i, column=3, value=people)
    wb.save(path)
    return path


def test_a_bill_after_midnight_belongs_to_the_night_before():
    """The POS day runs 06:00 to 06:00."""
    late = parse_closed_at("3/01/2026 1:24:20 a. m.")
    early = parse_closed_at("2/01/2026 8:00:00 p. m.")

    assert operating_day(late) == datetime.date(2026, 1, 2)
    assert operating_day(early) == datetime.date(2026, 1, 2)


def test_six_in_the_morning_starts_the_new_day():
    assert operating_day(parse_closed_at("3/01/2026 6:30:00 a. m.")) == datetime.date(
        2026, 1, 3
    )


def test_bills_group_into_operating_days(tmp_path):
    path = _export(tmp_path / "ventas.xlsx", [
        ("2/01/2026 8:00:00 p. m.", 4),
        ("3/01/2026 1:30:00 a. m.", 2),   # still the 2nd
        ("3/01/2026 9:00:00 p. m.", 3),
    ])

    days = read_daily_covers(path)

    assert [(d.date, d.guests, d.accounts) for d in days] == [
        (datetime.date(2026, 1, 2), 6, 2),
        (datetime.date(2026, 1, 3), 3, 1),
    ]


def test_a_file_without_the_people_column_is_refused(tmp_path):
    wb = Workbook()
    wb.active.cell(row=1, column=1, value="otra cosa")
    path = tmp_path / "otro.xlsx"
    wb.save(path)

    with pytest.raises(TicketExportError, match="NÚMERODEPERSONAS"):
        read_daily_covers(path)


def test_it_reads_the_file_by_content_not_by_extension(tmp_path):
    """The POS names it .XLS though it is really an .xlsx."""
    real = _export(tmp_path / "real.xlsx", [("2/01/2026 8:00:00 p. m.", 4)])
    mislabelled = tmp_path / "VENTAS.XLS"
    mislabelled.write_bytes(real.read_bytes())

    assert read_daily_covers(mislabelled)[0].guests == 4


@pytest.mark.django_db
def test_importing_twice_updates_instead_of_duplicating(tmp_path, restaurant):
    path = _export(tmp_path / "ventas.xlsx", [("2/01/2026 8:00:00 p. m.", 4)])
    call_command("import_covers", "--file", str(path),
                 "--restaurant", restaurant.slug, stdout=StringIO())

    # A corrected export for the same day.
    _export(path, [("2/01/2026 8:00:00 p. m.", 4), ("2/01/2026 9:00:00 p. m.", 5)])
    call_command("import_covers", "--file", str(path),
                 "--restaurant", restaurant.slug, stdout=StringIO())

    covers = DailyCovers.objects.get(restaurant=restaurant, date=datetime.date(2026, 1, 2))
    assert (covers.guests, covers.accounts) == (9, 2)
    assert DailyCovers.objects.count() == 1


@pytest.mark.django_db
def test_party_size(restaurant):
    covers = DailyCovers.objects.create(
        restaurant=restaurant, date=datetime.date(2026, 1, 2), guests=9, accounts=3
    )
    assert covers.party_size == 3

    empty = DailyCovers.objects.create(
        restaurant=restaurant, date=datetime.date(2026, 1, 3), guests=0, accounts=0
    )
    assert empty.party_size is None


def _workbook(path, rows):
    """The owner's master workbook: turnout lives in "Datos totales"."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos totales "  # the real sheet carries a trailing space
    ws.cell(row=1, column=2, value="Fecha")
    for i, (when, guests, accounts) in enumerate(rows, start=2):
        ws.cell(row=i, column=2, value=when)
        ws.cell(row=i, column=6, value=guests)
        ws.cell(row=i, column=7, value=accounts)
    wb.save(path)
    return path


def test_the_workbook_carries_the_history():
    """Four years of turnout live in the sheet, not in the POS export."""
    import tempfile

    from sales.importers.tickets import read_covers_from_workbook

    with tempfile.TemporaryDirectory() as tmp:
        path = _workbook(f"{tmp}/master.xlsx", [
            (datetime.datetime(2022, 1, 6), 25, 7),
            (datetime.datetime(2026, 9, 27), 48, 17),
        ])
        days = read_covers_from_workbook(path)

    assert [(d.date, d.guests, d.accounts) for d in days] == [
        (datetime.date(2022, 1, 6), 25, 7),
        (datetime.date(2026, 9, 27), 48, 17),
    ]


def test_a_day_with_no_turnout_stays_unknown():
    """An empty cell must not become "nobody came"."""
    import tempfile

    from sales.importers.tickets import read_covers_from_workbook

    with tempfile.TemporaryDirectory() as tmp:
        path = _workbook(f"{tmp}/master.xlsx", [
            (datetime.datetime(2022, 1, 6), None, None),
            (datetime.datetime(2022, 1, 7), 0, 0),
            (datetime.datetime(2022, 1, 8), 10, 4),
        ])
        days = read_covers_from_workbook(path)

    assert [d.date for d in days] == [datetime.date(2022, 1, 8)]
